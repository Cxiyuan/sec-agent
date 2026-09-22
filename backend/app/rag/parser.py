"""文档解析：txt / md / pdf / docx / xlsx -> 纯文本。"""

from __future__ import annotations

import io
from pathlib import Path

from app.config import settings

TEXT_EXT = (".txt", ".md", ".markdown", ".log", ".csv")


def decode_text(data: bytes) -> str:
    """按配置顺序探测编码解码（默认 utf-8 -> gb18030）；全部失败抛 ValueError。"""
    encodings = [e.strip() for e in settings.text_encodings.split(",") if e.strip()]
    for enc in encodings:
        try:
            text = data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        return text.lstrip("\ufeff")  # 去掉 BOM
    raise ValueError(f"无法识别的文本编码（已尝试：{', '.join(encodings)}）")


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in TEXT_EXT:
        return decode_text(data)
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext == ".xlsx":
        return _xlsx(data)
    # 兜底：按文本解码
    return decode_text(data)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
